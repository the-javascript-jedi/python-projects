import pandas as pd
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import PP_PLACEHOLDER


# -------------------------
# CONFIG
# -------------------------
EXCEL_FILE = "slides.xlsx"
PPT_TEMPLATE = "BaseTemplate.pptx"
OUTPUT_PPT = "FinalDeck.pptx"


# -------------------------
# HELPERS
# -------------------------
def get_placeholder(slide, placeholder_type):
    for shape in slide.placeholders:
        if shape.placeholder_format.type == placeholder_type:
            return shape
    return None


# -------------------------
# MAIN
# -------------------------
def main():
    prs = Presentation(PPT_TEMPLATE)
    df = pd.read_excel(EXCEL_FILE)

    layout_map = {
        "intro": prs.slide_layouts[0],     # Title Slide
        "feature": prs.slide_layouts[1],   # Title + Content
        "summary": prs.slide_layouts[2],   # Section Header
    }

    for _, row in df.iterrows():
        slide_type = row["slide_type"]
        layout = layout_map.get(slide_type, prs.slide_layouts[1])
        slide = prs.slides.add_slide(layout)

        # -------------------------
        # TITLE
        # -------------------------
        if slide.shapes.title and pd.notna(row["title"]):
            slide.shapes.title.text = str(row["title"])

        # -------------------------
        # TEXT CONTENT
        # -------------------------
        text_parts = []

        if pd.notna(row["subtitle"]):
            text_parts.append(str(row["subtitle"]))

        if pd.notna(row["content"]):
            text_parts.append(str(row["content"]))

        full_text = "\n".join(text_parts)

        body = get_placeholder(slide, PP_PLACEHOLDER.BODY)
        subtitle = get_placeholder(slide, PP_PLACEHOLDER.SUBTITLE)

        if body:
            body.text = full_text
        elif subtitle:
            subtitle.text = full_text

        # -------------------------
        # IMAGE (optional)
        # -------------------------
        if pd.notna(row["image_path"]):
            try:
                slide.shapes.add_picture(
                    row["image_path"],
                    Inches(5),
                    Inches(2),
                    width=Inches(4)
                )
            except Exception as e:
                print(f"⚠️ Image skipped: {e}")

    prs.save(OUTPUT_PPT)
    print(f"✅ Presentation created: {OUTPUT_PPT}")


if __name__ == "__main__":
    main()
